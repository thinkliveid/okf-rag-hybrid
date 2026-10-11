from __future__ import annotations

import csv
import json
from io import StringIO
from pathlib import Path

from .models import FileInfo, ParsedDocument


_SPLIT_HEADINGS: dict[str, list[str]] = {}



def _split_by_product_qa(text: str) -> list[tuple[str, str]]:
    lines = text.splitlines()
    blocks = []
    current_block = []
    current_title = ""
    for i, line in enumerate(lines):
        if line.strip().startswith("Product:"):
            if current_block:
                if len(current_block) >= 1 and current_block[-1].strip() and not current_block[-1].strip().startswith("Product:"):
                    title_candidate = current_block.pop().strip()
                    if current_block:
                        blocks.append((current_title or title_candidate, "\n".join(current_block).strip()))
                    current_title = title_candidate
                    current_block = [title_candidate, line]
                else:
                    blocks.append((current_title or "General Product Info", "\n".join(current_block).strip()))
                    current_title = line.split(":", 1)[1].strip()
                    current_block = [line]
            else:
                current_title = line.split(":", 1)[1].strip()
                current_block = [line]
        else:
            current_block.append(line)
    if current_block:
        blocks.append((current_title or "General Product Info", "\n".join(current_block).strip()))
    return [b for b in blocks if b[1]]


def _split_by_headings(text: str, headings: list[str]) -> list[tuple[str, str]]:
    lines = text.splitlines()
    blocks = []
    current_block = []
    current_title = headings[0]
    heading_set = set(h.strip().lower() for h in headings)

    for line in lines:
        stripped = line.strip()
        if stripped.lower() in heading_set:
            if current_block:
                blocks.append((current_title, "\n".join(current_block).strip()))
            exact_title = next(h for h in headings if h.strip().lower() == stripped.lower())
            current_title = exact_title
            current_block = [line]
        else:
            current_block.append(line)

    if current_block:
        blocks.append((current_title, "\n".join(current_block).strip()))
    return [b for b in blocks if b[1]]


def parse(files: list[FileInfo]) -> list[ParsedDocument]:
    parsed_docs = []
    for file_info in files:
        parsed_docs.extend(parse_file(file_info))
    return parsed_docs


def parse_file(file_info: FileInfo) -> list[ParsedDocument]:
    ext = file_info.extension
    path = file_info.path
    metadata_base = {
        "source": str(path),
        "extension": ext,
        "mime_type": file_info.mime_type,
        "size": file_info.size,
        "modified_at": file_info.modified_at.isoformat(),
    }

    if ext in {"txt", "md", "markdown", "html", "htm", "log", "sql", "yaml", "yml"}:
        text = path.read_text(encoding="utf-8", errors="ignore")
    elif ext == "json":
        text = _parse_json(path)
    elif ext == "csv":
        text = _parse_csv(path)
    elif ext == "docx":
        text = _parse_docx(path)
    elif ext == "pdf":
        text = _parse_pdf(path)
    else:
        text = path.read_text(encoding="utf-8", errors="ignore")

    text = text.strip()
    filename = path.name

    if "Product:" in text:
        segments = _split_by_product_qa(text)
    elif filename in _SPLIT_HEADINGS:
        segments = _split_by_headings(text, _SPLIT_HEADINGS[filename])
    else:
        segments = [(file_info.path.stem.replace("_", " ").replace("-", " ").title(), text)]

    docs = []
    seen_titles = {}
    for title, content in segments:
        if title in seen_titles:
            seen_titles[title] += 1
            unique_title = f"{title} {seen_titles[title]}"
        else:
            seen_titles[title] = 1
            unique_title = title

        metadata = dict(metadata_base)
        metadata["segment_title"] = unique_title
        metadata["source"] = f"{path}#{unique_title}"
        docs.append(ParsedDocument(file=file_info, text=content, metadata=metadata))
    return docs


def _parse_json(path: Path) -> str:
    data = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
    return json.dumps(data, indent=2, ensure_ascii=False)


def _parse_csv(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    reader = csv.DictReader(StringIO(raw))
    rows = list(reader)
    if rows:
        return "\n".join(json.dumps(row, ensure_ascii=False) for row in rows)
    return raw


def _parse_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Install python-docx to parse .docx files") from exc
    doc = Document(str(path))
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _parse_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("Install pypdf to parse .pdf files") from exc
    reader = PdfReader(str(path))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)
